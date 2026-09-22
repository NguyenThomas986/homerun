"""Create a project-level Excel workbook containing machine-readable QC.

The workbook complements the HTML/PDF report. It keeps one row per replicate
so lane-merged tag directories and replicate-specific peak calls stay visible.
"""
from __future__ import annotations

from collections import Counter
import csv
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

from .utils import assay_of_leaf, iter_leaf_dirs, log, replicate_of_leaf


HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)
TITLE_FILL = PatternFill("solid", fgColor="D9EAF7")


def _taginfo_value(lines, key):
    for line in lines:
        if line.startswith(key):
            return line.split("=")[-1].strip()
    return "NA"


def _tagdir_rows(cfg):
    source_counts = Counter()
    identities = set()

    for species, sample, leaf_name, _r1 in iter_leaf_dirs(cfg):
        key = (species, sample, leaf_name)
        source_counts[key] += 1
        identities.add(key)

    rows = []

    for species, sample, leaf_name in sorted(identities):
        tagdir = cfg.leaf_tagdir(species, sample, leaf_name)

        if not tagdir.is_dir():
            continue

        row = {
            "Species": species,
            "Sample": sample,
            "Assay": assay_of_leaf(leaf_name) or "NA",
            "Replicate": replicate_of_leaf(leaf_name) or "NA",
            "Library": leaf_name,
            "Source files": source_counts[(species, sample, leaf_name)],
            "TagDir": str(tagdir),
            "Total Tags": "NA",
            "Unique Positions": "NA",
            "Tags per BP": "NA",
            "Average Tags/Position": "NA",
            "Median Tags/Position": "NA",
            "Average Read Length": "NA",
            "Average Fragment GC": "NA",
        }

        info = tagdir / "tagInfo.txt"

        if info.exists():
            lines = info.read_text(errors="replace").splitlines()
            genome_line = next(
                (line for line in lines if line.startswith("genome=")),
                "",
            )
            parts = genome_line.split("\t")

            row.update(
                {
                    "Total Tags": (
                        parts[2].strip() if len(parts) > 2 else "NA"
                    ),
                    "Unique Positions": (
                        parts[1].strip() if len(parts) > 1 else "NA"
                    ),
                    "Tags per BP": _taginfo_value(lines, "tagsPerBP"),
                    "Average Tags/Position": _taginfo_value(
                        lines,
                        "averageTagsPerPosition",
                    ),
                    "Median Tags/Position": _taginfo_value(
                        lines,
                        "medianTagsPerPosition",
                    ),
                    "Average Read Length": _taginfo_value(
                        lines,
                        "averageTagLength",
                    ),
                    "Average Fragment GC": _taginfo_value(
                        lines,
                        "averageFragmentGCcontent",
                    ),
                }
            )

        rows.append(row)

    return rows


def _grab(text, pattern, cast=str, default="NA"):
    match = re.search(pattern, text, flags=re.MULTILINE)

    if not match:
        return default

    value = match.group(1).strip().rstrip("%")

    try:
        return cast(value)
    except (TypeError, ValueError):
        return default


def _grab_any(text, patterns, cast=str, default="NA"):
    """Try multiple possible HOMER labels for the same statistic."""
    for pattern in patterns:
        value = _grab(
            text,
            pattern,
            cast=cast,
            default=None,
        )

        if value is not None:
            return value

    return default


def _tss_annotation_counts(tss_file):
    """Count TSR annotation categories from a matching HOMER *.tss.txt file."""
    result = {
        "TSRs w. annotation": "NA",
        "tss": "NA",
        "first Exon": "NA",
        "single Exon": "NA",
        "tssAntisense": "NA",
        "exon": "NA",
        "other": "NA",
    }

    if not tss_file.exists():
        return result

    try:
        with tss_file.open(newline="", errors="replace") as handle:
            reader = csv.DictReader(handle, delimiter="\t")
            rows = list(reader)
    except (OSError, csv.Error):
        return result

    if not rows:
        return result

    annotation_counts = Counter()

    for row in rows:
        # HOMER TSS files can include non-data/header-like rows. Match qc.py's
        # behavior by only keeping chromosome rows when a chr field exists.
        chrom = row.get("chr")
        if chrom is not None and not chrom.startswith("chr"):
            continue

        annotation = (row.get("annotation") or "").strip()

        if annotation:
            annotation_counts[annotation] += 1

    valid_total = sum(annotation_counts.values())

    result["TSRs w. annotation"] = valid_total
    result["tss"] = annotation_counts.get("tss", 0)
    result["first Exon"] = annotation_counts.get("firstExon", 0)
    result["single Exon"] = annotation_counts.get("singleExon", 0)
    result["tssAntisense"] = annotation_counts.get("tssAntisense", 0)

    # Combine exon-style categories into the single requested "exon" column.
    result["exon"] = (
        annotation_counts.get("otherExon", 0)
        + annotation_counts.get("otherExonBidirectional", 0)
    )

    result["other"] = annotation_counts.get("other", 0)

    return result


def _peak_rows(cfg):
    """Return one QC row per findcsRNATSS peak call."""
    rows = []

    for species_dir in sorted(
        path for path in cfg.project.iterdir()
        if path.is_dir()
    ):
        tss_dir = species_dir / "TSS"

        if not tss_dir.is_dir():
            continue

        for stats in sorted(tss_dir.glob("*.stats.txt")):
            text = stats.read_text(errors="replace")
            library = stats.name.removesuffix(".stats.txt")
            sample = library.split("_", 1)[0]
            tss_file = tss_dir / f"{library}.tss.txt"

            annotations = _tss_annotation_counts(tss_file)

            stable_tss = _grab_any(
                text,
                [
                    r"Stable TSSs?:\s+(\d+)",
                    r"stableTSSs\s*[:=]\s*(\d+)",
                    r"Stable TSS count:\s+(\d+)",
                ],
                int,
            )

            unstable_tss = _grab_any(
                text,
                [
                    r"Unstable TSSs?:\s+(\d+)",
                    r"unstableTSSs\s*[:=]\s*(\d+)",
                    r"Unstable TSS count:\s+(\d+)",
                ],
                int,
            )

            total_tss = "NA"
            pct_stable = "NA"
            pct_unstable = "NA"

            if isinstance(stable_tss, int) and isinstance(unstable_tss, int):
                total_tss = stable_tss + unstable_tss

                if total_tss > 0:
                    pct_stable = round(
                        100.0 * stable_tss / total_tss,
                        2,
                    )
                    pct_unstable = round(
                        100.0 * unstable_tss / total_tss,
                        2,
                    )

            putative_tsrs = _grab(
                text,
                r"total putative TSS clusters\s+(\d+)",
                int,
            )

            valid_tsrs = _grab(
                text,
                r"Valid TSS clusters\s+(\d+)",
                int,
            )

            stable_tsrs = _grab_any(
                text,
                [
                    r"Stable TSRs?:\s+(\d+)",
                    r"stable TSRs?:\s+(\d+)",
                    r"stable TSS clusters\s+(\d+)",
                ],
                int,
            )

            unstable_tsrs = _grab_any(
                text,
                [
                    r"Unstable TSRs?:\s+(\d+)",
                    r"unstable TSRs?:\s+(\d+)",
                    r"unstable TSS clusters\s+(\d+)",
                ],
                int,
            )

            row = {
                "Species": species_dir.name,
                "Sample": sample,
                "Peak call": library,

                "Total csRNA reads": _grab(
                    text,
                    r"Total csRNA reads:\s+([\d.]+)",
                    float,
                ),
                "Total input reads": _grab(
                    text,
                    r"Total input reads:\s+([\d.]+)",
                    float,
                ),

                # TSS stability counts
                "stableTSSs": stable_tss,
                "unstableTSSs": unstable_tss,
                "TSS": total_tss,
                "% unstable": pct_unstable,
                "% stable": pct_stable,

                # TSR counts
                "putative TSRs": putative_tsrs,
                "valid TSRs": valid_tsrs,
                "stableTSRs": stable_tsrs,
                "unstableTSRs": unstable_tsrs,

                # TSR percentages
                "Bidirectional TSRs [%]": _grab(
                    text,
                    r"Fraction of bidirectional.*?:\s+([\d.]+%)",
                    float,
                ),
                "Stable TSRs [%]": _grab(
                    text,
                    r"Fraction of stable.*?:\s+([\d.]+%)",
                    float,
                ),

                # Stability classes
                "S": _grab_any(
                    text,
                    [
                        r"^\s*S:\s+\d+\s+\(([\d.]+%)",
                    ],
                    float,
                ),
                "SS": _grab(
                    text,
                    r"^\s*SS:\s+\d+\s+\(([\d.]+%)",
                    float,
                ),
                "SU": _grab(
                    text,
                    r"^\s*SU:\s+\d+\s+\(([\d.]+%)",
                    float,
                ),
                "U": _grab_any(
                    text,
                    [
                        r"^\s*U:\s+\d+\s+\(([\d.]+%)",
                    ],
                    float,
                ),
                "US": _grab(
                    text,
                    r"^\s*US:\s+\d+\s+\(([\d.]+%)",
                    float,
                ),
                "UU": _grab(
                    text,
                    r"^\s*UU:\s+\d+\s+\(([\d.]+%)",
                    float,
                ),

                # Annotation counts from *.tss.txt
                **annotations,

                # Existing useful stats
                "Distal %": _grab(
                    text,
                    r"Fraction Promoter-Distal.*?:\s+([\d.]+%)",
                    float,
                ),
                "Log2 vs Input": _grab(
                    text,
                    r"log2 fold vs\. input:\s+([\d.\-]+)",
                    float,
                ),
                "Log2 vs RNA": _grab(
                    text,
                    r"log2 fold vs\. rna:\s+([\d.\-]+)",
                    float,
                ),

                "Stats file": str(stats),
                "TSS file": str(tss_file),
            }

            rows.append(row)

    return rows


def _parse_lengths(path):
    """Return total, retained, adapter/dimer and weighted retained length."""
    try:
        with path.open(newline="") as handle:
            reader = csv.reader(handle, delimiter="\t")
            next(reader, None)

            values = []

            for row in reader:
                if len(row) < 2:
                    continue

                try:
                    values.append(
                        (
                            float(row[0]),
                            float(row[1]),
                        )
                    )
                except ValueError:
                    continue

    except OSError:
        return None

    if not values:
        return None

    total = sum(count for _length, count in values)

    if total <= 0:
        return None

    adapters = sum(
        count
        for length, count in values
        if length == 0
    )

    retained = total - adapters

    average = (
        sum(
            length * count
            for length, count in values
            if length != 0
        )
        / retained
        if retained > 0
        else "NA"
    )

    return {
        "Input Reads": int(total),
        "Retained Reads": int(retained),
        "Retained %": round(100 * retained / total, 2),
        "Adapter/Dimer Reads": int(adapters),
        "Adapter/Dimer %": round(100 * adapters / total, 2),
        "Average Retained Length": (
            round(average, 2)
            if average != "NA"
            else "NA"
        ),
    }


def _parse_star(path):
    text = path.read_text(errors="replace")

    patterns = {
        "Input Reads": (
            r"Number of input reads \|\s+(\d+)",
            int,
        ),
        "Uniquely Mapped %": (
            r"Uniquely mapped reads % \|\s+([\d.]+)%",
            float,
        ),
        "Multi-Mapped %": (
            r"% of reads mapped to multiple loci \|\s+([\d.]+)%",
            float,
        ),
        "Too-Many-Loci %": (
            r"% of reads mapped to too many loci \|\s+([\d.]+)%",
            float,
        ),
        "Unmapped (mismatch) %": (
            r"% of reads unmapped: too many mismatches \|\s+([\d.]+)%",
            float,
        ),
        "Unmapped (short) %": (
            r"% of reads unmapped: too short \|\s+([\d.]+)%",
            float,
        ),
        "Unmapped (other) %": (
            r"% of reads unmapped: other \|\s+([\d.]+)%",
            float,
        ),
    }

    result = {}

    for name, (pattern, cast) in patterns.items():
        match = re.search(pattern, text)

        if match:
            result[name] = cast(match.group(1))

    unmapped = [
        value
        for key, value in result.items()
        if key.startswith("Unmapped (")
    ]

    if unmapped:
        result["Unmapped %"] = round(sum(unmapped), 2)

    uniquely = result.get("Uniquely Mapped %")
    multi = result.get("Multi-Mapped %")
    too_many = result.get("Too-Many-Loci %")

    if all(
        value is not None
        for value in (
            uniquely,
            multi,
            too_many,
        )
    ):
        result["Overall Aligned %"] = round(
            uniquely + multi + too_many,
            2,
        )

    return result


def _parse_hisat2(path):
    text = path.read_text(errors="replace")
    result = {}

    patterns = {
        "Input Reads": (
            r"(\d+)\s+reads?;\s+of these",
            int,
        ),
        "Unmapped %": (
            r"([\d.]+)\s*%\)\s+aligned 0 times",
            float,
        ),
        "Uniquely Mapped %": (
            r"([\d.]+)\s*%\)\s+aligned exactly 1 time",
            float,
        ),
        "Multi-Mapped %": (
            r"([\d.]+)\s*%\)\s+aligned >1 times",
            float,
        ),
        "Overall Aligned %": (
            r"([\d.]+)\s*%\s+overall alignment rate",
            float,
        ),
    }

    for name, (pattern, cast) in patterns.items():
        match = re.search(pattern, text)

        if match:
            result[name] = cast(match.group(1))

    return result


def _trim_alignment_rows(cfg):
    trim_rows = []
    alignment_rows = []

    for species, sample, leaf_name, r1 in iter_leaf_dirs(cfg):
        assay = assay_of_leaf(leaf_name) or "NA"
        replicate = replicate_of_leaf(leaf_name) or "NA"

        common = {
            "Species": species,
            "Sample": sample,
            "Assay": assay,
            "Replicate": replicate,
            "Library": leaf_name,
            "FASTQ": r1.name,
        }

        lengths = (
            cfg.trimmed_dir(species, sample)
            / f"{r1.name}.lengths"
        )

        if lengths.exists():
            parsed = _parse_lengths(lengths)

            if parsed:
                trim_rows.append(
                    {
                        **common,
                        "Tool": "homerTools",
                        **parsed,
                        "Source log": str(lengths),
                    }
                )

        prefix = r1.name.split("_R1")[0]
        aligned_dir = cfg.aligned_dir(species, sample)

        star_log = aligned_dir / f"{prefix}.Log.final.out"
        hisat2_log = aligned_dir / f"{prefix}_mappingstats.txt"

        if star_log.exists():
            parsed = _parse_star(star_log)
            tool = "STAR"
            source = star_log

        elif hisat2_log.exists():
            parsed = _parse_hisat2(hisat2_log)
            tool = "HISAT2"
            source = hisat2_log

        else:
            parsed = None

        if parsed:
            alignment_rows.append(
                {
                    **common,
                    "Aligner": tool,
                    "Input Reads": parsed.get(
                        "Input Reads",
                        "NA",
                    ),
                    "Uniquely Mapped %": parsed.get(
                        "Uniquely Mapped %",
                        "NA",
                    ),
                    "Multi-Mapped %": parsed.get(
                        "Multi-Mapped %",
                        "NA",
                    ),
                    "Unmapped %": parsed.get(
                        "Unmapped %",
                        "NA",
                    ),
                    "Overall Aligned %": parsed.get(
                        "Overall Aligned %",
                        "NA",
                    ),
                    "Source log": str(source),
                }
            )

    return trim_rows, alignment_rows


def _write_table_sheet(
    workbook,
    name,
    rows,
    table_name,
):
    ws = workbook.create_sheet(name)

    if not rows:
        ws.append(
            [
                f"No {name.lower()} data were found."
            ]
        )

        ws["A1"].fill = TITLE_FILL
        ws["A1"].font = Font(bold=True)

        return ws

    # Use every key present anywhere in the rows while preserving first-seen
    # column order. This keeps optional TSR fields from disappearing if the
    # first row happens to contain fewer values.
    headers = []

    for row in rows:
        for key in row:
            if key not in headers:
                headers.append(key)

    ws.append(headers)

    for row in rows:
        ws.append(
            [
                row.get(header, "NA")
                for header in headers
            ]
        )

    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    table = Table(
        displayName=table_name,
        ref=ws.dimensions,
    )

    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )

    ws.add_table(table)

    for column_cells in ws.columns:
        values = [
            str(cell.value or "")
            for cell in column_cells
        ]

        width = min(
            max(
                max(map(len, values)) + 2,
                11,
            ),
            45,
        )

        ws.column_dimensions[
            column_cells[0].column_letter
        ].width = width

    # Format all percentage-style columns numerically.
    for cell in ws[1]:
        header = str(cell.value or "")

        if (
            "%" in header
            or header.startswith("%")
            or header.endswith("[%]")
        ):
            for data_cell in ws.iter_cols(
                min_col=cell.column,
                max_col=cell.column,
                min_row=2,
            ):
                data_cell[0].number_format = "0.00"

    return ws


def _write_overview(
    workbook,
    cfg,
    tag_rows,
    peak_rows,
    trim_rows,
    alignment_rows,
):
    ws = workbook.create_sheet("Overview")

    ws.append(["HOMERun QC Summary"])

    ws["A1"].fill = HEADER_FILL
    ws["A1"].font = Font(
        color="FFFFFF",
        bold=True,
        size=14,
    )

    ws.merge_cells("A1:B1")

    ws.append(
        [
            "Project",
            str(cfg.project),
        ]
    )
    ws.append(
        [
            "Replicate TagDirs",
            len(tag_rows),
        ]
    )
    ws.append(
        [
            "Peak calls",
            len(peak_rows),
        ]
    )
    ws.append(
        [
            "Trimming records",
            len(trim_rows),
        ]
    )
    ws.append(
        [
            "Alignment records",
            len(alignment_rows),
        ]
    )
    ws.append(
        [
            "Species",
            len(
                {
                    row["Species"]
                    for row in tag_rows
                }
            ),
        ]
    )
    ws.append(
        [
            "Samples",
            len(
                {
                    (
                        row["Species"],
                        row["Sample"],
                    )
                    for row in tag_rows
                }
            ),
        ]
    )

    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 70

    for row in range(
        2,
        ws.max_row + 1,
    ):
        ws.cell(
            row,
            1,
        ).font = Font(bold=True)


def write_qc_workbook(cfg):
    """Write ``QC_summary.xlsx`` at the project root and return its path."""
    tag_rows = _tagdir_rows(cfg)
    peak_rows = _peak_rows(cfg)
    trim_rows, alignment_rows = _trim_alignment_rows(cfg)

    workbook = Workbook()
    workbook.remove(workbook.active)

    _write_overview(
        workbook,
        cfg,
        tag_rows,
        peak_rows,
        trim_rows,
        alignment_rows,
    )

    _write_table_sheet(
        workbook,
        "Trimming",
        trim_rows,
        "TrimmingQC",
    )

    _write_table_sheet(
        workbook,
        "Alignment",
        alignment_rows,
        "AlignmentQC",
    )

    _write_table_sheet(
        workbook,
        "Tag Directories",
        tag_rows,
        "TagDirectoryQC",
    )

    _write_table_sheet(
        workbook,
        "Peak Calling",
        peak_rows,
        "PeakCallingQC",
    )

    output = cfg.project / "QC_summary.xlsx"

    workbook.save(output)

    log.info(
        "QC: wrote Excel workbook %s",
        output,
    )

    return output
