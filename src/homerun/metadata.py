"""Load and validate optional CSV/XLSX FASTQ metadata manifests."""
from __future__ import annotations

import csv
import glob
import re
from functools import lru_cache
from pathlib import Path
from typing import Iterable


REQUIRED_COLUMNS = {"fastq", "species", "assay", "replicate"}
OPTIONAL_COLUMNS = {"sample", "condition"}
SUPPORTED_ASSAYS = {
    "csrna": "csRNA",
    "srna": "sRNA",
    "totalrna": "totalRNA",
    "rna": "RNA",
}


def _column_key(value) -> str:
    """Normalize headings case-insensitively and ignore spaces/_/- chars."""
    return re.sub(r"[\s_-]+", "", str(value or "").strip()).lower()


def _cell(value) -> str:
    return "" if value is None else str(value).strip()


def _csv_rows(path: Path) -> tuple[list[str], Iterable[tuple[int, list]]]:
    handle = path.open(newline="", encoding="utf-8-sig")
    reader = csv.reader(handle)
    try:
        headers = next(reader)
    except StopIteration:
        handle.close()
        return [], ()

    def rows():
        try:
            for row_number, row in enumerate(reader, start=2):
                yield row_number, row
        finally:
            handle.close()

    return headers, rows()


def _xlsx_rows(path: Path) -> tuple[list[str], Iterable[tuple[int, list]]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - declared dependency
        raise ValueError(
            "Reading .xlsx metadata requires openpyxl; install HomeRun's "
            "declared dependencies and try again."
        ) from exc

    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    iterator = sheet.iter_rows(values_only=True)
    try:
        headers = list(next(iterator))
    except StopIteration:
        workbook.close()
        return [], ()

    def rows():
        try:
            for row_number, row in enumerate(iterator, start=2):
                yield row_number, list(row)
        finally:
            workbook.close()

    return headers, rows()


def _read_rows(path: Path) -> tuple[list[str], Iterable[tuple[int, list]]]:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        return _csv_rows(path)
    if suffix == ".xlsx":
        return _xlsx_rows(path)
    raise ValueError(
        f"Unsupported metadata extension '{path.suffix}' for '{path}'. "
        "Use a .csv or .xlsx file."
    )


def _read_role(filename: str) -> str | None:
    """Return R1/R2 when the filename has a conventional read marker."""
    match = re.search(r"(?:^|[-_])(R[12])(?=[-_.]|$)", filename)
    return match.group(1) if match else None


def load_sample_metadata(path) -> dict[str, dict[str, str]]:
    """Return validated sample metadata keyed by exact FASTQ basename.

    Column headings are case-insensitive and tolerate spaces, underscores,
    and hyphens. Species and replicate values are normalized to lowercase;
    assay spelling is normalized to HomeRun's existing canonical tokens.
    """
    metadata_path = Path(path).expanduser()
    if not metadata_path.is_file():
        raise ValueError(f"Metadata file does not exist: {metadata_path}")

    headers, rows = _read_rows(metadata_path)
    normalized = [_column_key(header) for header in headers]
    if not normalized or not any(normalized):
        raise ValueError(f"Metadata file has no header row: {metadata_path}")

    duplicates = sorted({name for name in normalized if name and normalized.count(name) > 1})
    if duplicates:
        raise ValueError(
            "Metadata contains duplicate columns after normalization: "
            + ", ".join(duplicates)
        )

    missing = sorted(REQUIRED_COLUMNS - set(normalized))
    if missing:
        display = {"fastq": "FASTQ", "species": "Species", "assay": "Assay", "replicate": "Replicate"}
        raise ValueError(
            "Metadata is missing required column(s): "
            + ", ".join(display[name] for name in missing)
        )

    indexes = {name: normalized.index(name) for name in REQUIRED_COLUMNS | OPTIONAL_COLUMNS if name in normalized}
    metadata: dict[str, dict[str, str]] = {}
    identities: dict[tuple[str, str, str], list[str]] = {}

    for row_number, values in rows:
        if not any(_cell(value) for value in values):
            continue

        def value(column: str) -> str:
            index = indexes.get(column)
            return _cell(values[index]) if index is not None and index < len(values) else ""

        fastq = value("fastq")
        species = value("species").lower()
        assay_input = value("assay")
        replicate = value("replicate").lower()

        for column, field_value in (
            ("FASTQ", fastq),
            ("Species", species),
            ("Assay", assay_input),
            ("Replicate", replicate),
        ):
            if not field_value:
                raise ValueError(f"Metadata row {row_number} has blank required field {column}")

        if "/" in fastq or "\\" in fastq:
            raise ValueError(
                f"Metadata row {row_number} FASTQ must be a filename, not a path: '{fastq}'"
            )
        if fastq in metadata:
            raise ValueError(f"Duplicate FASTQ metadata entry at row {row_number}: '{fastq}'")

        assay = SUPPORTED_ASSAYS.get(assay_input.lower())
        if assay is None:
            supported = ", ".join(SUPPORTED_ASSAYS.values())
            raise ValueError(
                f"Invalid assay '{assay_input}' at metadata row {row_number}; "
                f"supported assays are: {supported}"
            )
        if not re.fullmatch(r"r(?:ep)?\d+", replicate):
            raise ValueError(
                f"Invalid replicate '{value('replicate')}' at metadata row {row_number}; "
                "use forms such as r1, r2, rep1, or rep2"
            )

        sample = value("sample") or species
        condition = value("condition")
        for label, component in (("Species", species), ("Sample", sample), ("Condition", condition)):
            if component and ("/" in component or "\\" in component):
                raise ValueError(
                    f"Metadata row {row_number} {label} cannot contain '/' or '\\': '{component}'"
                )

        leaf_name = "_".join(part for part in (condition, assay, replicate) if part)
        record = {
            "species": species,
            "sample": sample,
            "condition": condition,
            "assay": assay,
            "replicate": replicate,
            "leaf_name": leaf_name,
        }
        metadata[fastq] = record
        identities.setdefault((species, sample, leaf_name), []).append(fastq)

    if not metadata:
        raise ValueError(f"Metadata file contains no sample rows: {metadata_path}")

    for (species, sample, leaf_name), filenames in identities.items():
        if len(filenames) < 2:
            continue
        roles = [_read_role(filename) for filename in filenames]
        if len(filenames) == 2 and set(roles) == {"R1", "R2"}:
            continue
        raise ValueError(
            "Duplicate metadata identity would collide for "
            f"{species}/{sample}/{leaf_name}: {', '.join(filenames)}"
        )

    return metadata


@lru_cache(maxsize=16)
def _cached_metadata(path: str, modified_ns: int, size: int):
    """Cache parsed manifests while invalidating after an on-disk change."""
    del modified_ns, size
    return load_sample_metadata(path)


def metadata_for_config(cfg) -> dict[str, dict[str, str]] | None:
    """Load cfg.metadata once, or return None when filename mode is active."""
    configured = getattr(cfg, "metadata", "")
    if not configured:
        return None
    path = Path(configured).expanduser().resolve()
    if not path.is_file():
        raise ValueError(f"Metadata file does not exist: {path}")
    stat = path.stat()
    return _cached_metadata(str(path), stat.st_mtime_ns, stat.st_size)


def metadata_record_for_fastq(cfg, fastq) -> dict[str, str] | None:
    """Return an exact-basename metadata match, strictly when configured."""
    metadata = metadata_for_config(cfg)
    if metadata is None:
        return None
    name = Path(fastq).name
    try:
        return metadata[name]
    except KeyError as exc:
        raise ValueError(
            f"Discovered FASTQ '{name}' has no entry in metadata file "
            f"'{Path(cfg.metadata)}'"
        ) from exc


def _is_fastq(path: Path) -> bool:
    return path.is_file() and (
        path.name.endswith(".fastq") or path.name.endswith(".fastq.gz")
    )


def discovered_fastqs(cfg) -> list[Path]:
    """Find the FASTQs HomeRun can stage or already has under RawData/."""
    found = [path for path in cfg.project.glob("*") if _is_fastq(path)]
    found.extend(
        path for path in cfg.project.glob("*/RawData/*") if _is_fastq(path)
    )
    if getattr(cfg, "copy_src", ""):
        found.extend(
            path for value in glob.glob(cfg.copy_src)
            if _is_fastq(path := Path(value))
        )
    return sorted(set(found))


def validate_metadata_fastqs(cfg) -> None:
    """Require a one-to-one filename inventory when metadata mode is active."""
    metadata = metadata_for_config(cfg)
    if metadata is None:
        return

    discovered = discovered_fastqs(cfg)
    discovered_names = {path.name for path in discovered}
    missing_rows = sorted(discovered_names - set(metadata))
    if missing_rows:
        raise ValueError(
            "Discovered FASTQ file(s) missing from metadata: "
            + ", ".join(missing_rows)
        )

    missing_files = sorted(set(metadata) - discovered_names)
    if missing_files:
        raise ValueError(
            "Metadata FASTQ entr"
            + ("ies do" if len(missing_files) != 1 else "y does")
            + " not match any discovered FASTQ file: "
            + ", ".join(missing_files)
        )
